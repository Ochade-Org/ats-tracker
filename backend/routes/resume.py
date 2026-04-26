"""
Resume processing, generation, and matching for ATS Matcher Backend (PostgreSQL)
"""

import json
import os
import datetime
import PyPDF2
from io import BytesIO
from docx import Document
from flask import jsonify, g, request, send_file
from db.database import get_db
from routes.usage import check_usage_limit, record_usage
from config import MAX_SAVED_RESUMES, MAX_BATCH_RESUMES
from openai import OpenAI, AsyncOpenAI
from agents import Agent, Runner, trace, function_tool, OpenAIChatCompletionsModel


# ---------------------------
# LLM INITIALIZATION
# ---------------------------


try:
    api_key = os.getenv("OPENROUTER_API_KEY")

    
    
    if api_key:
        model = OpenAI(
            base_url="https://openrouter.ai/api/v1",
            api_key=api_key
        )
        # client = AsyncOpenAI(base_url="https://openrouter.ai/api/v1", api_key=api_key)

        # model = OpenAIChatCompletionsModel(model="deepseek-chat", openai_client=client)
    else:
        model = None

except Exception as e:
    print(f"ERROR initializing LLM client: {e}")
    model = None


# ---------------------------
# LLM CALL WRAPPER
# ---------------------------
async def llm_call(prompt):
    if not model:
        raise RuntimeError("LLM model not initialized")
    
    message=[{"role": "user", "content": prompt}]

    response = model.chat.completions.create(
        messages=[{"role": "user", "content": prompt}],
        model="gpt-5-nano"
    )

    return response.choices[0].message.content

    # with trace("Protected Automated SDR"):
    #     result = await Runner.run(model, message)
    #     return result.final_output


# ---------------------------
# PDF TEXT EXTRACTION
# ---------------------------
def extract_text_from_pdf(pdf_stream):
    try:
        reader = PyPDF2.PdfReader(pdf_stream)
        text = ""

        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"

        return text

    except Exception as e:
        print(f"PDF extraction error: {e}")
        return None


# ---------------------------
# STANDARD RESUME GENERATION
# ---------------------------
def generate_standard_resume_pdf(resume_text):
    parse_prompt = f"""
    Parse the following resume text into structured JSON format for a professional CV.
    
    Resume Text:
    {resume_text}
    
    Extract and structure the information into the following sections:
    - name: Full name of the person (string)
    - contact: Object with email, phone, location, linkedin (strings, use empty string if not found)
    - summary: Professional summary or objective (string, use empty string if not present)
    - experience: Array of work experience objects, each with: title (string), company (string), duration (string), location (string), description (array of strings)
    - education: Array of education objects, each with: degree (string), institution (string), year (string), gpa (string)
    - skills: Array of technical/professional skills (strings only)
    - certifications: Array of certification names (strings only, no JSON formatting)
    - projects: Array of project objects, each with: name (string), description (string or array of strings), technologies (array of strings)
    
    IMPORTANT: 
    - Return ONLY valid JSON. Do not add any other text.
    - All string values should be plain text without quotes, brackets, or JSON formatting.
    - Arrays should contain only the actual content strings.
    - If a section is not present, use empty array [] or empty string "".
    - Do not include any JSON-like formatting in the string values themselves.
    """


    try:
        response_text = llm_call(parse_prompt)

        json_string = response_text.strip().replace("```json", "").replace("```", "")
        start = json_string.find("{")
        end = json_string.rfind("}") + 1
        json_string = json_string[start:end]

        parsed_data = json.loads(json_string)

    except Exception as e:
        print(f"Parsing error: {e}")
        parsed_data = {
            "name": "Professional Name",
            "contact": {},
            "summary": "",
            "experience": [],
            "education": [],
            "skills": [],
            "certifications": [],
            "projects": []
        }

    # ---------------------------
    # PDF GENERATION
    # ---------------------------
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib import colors

    pdf_stream = BytesIO()
    doc = SimpleDocTemplate(pdf_stream, pagesize=letter)

    styles = getSampleStyleSheet()
    story = []

    # Header
    story.append(Paragraph(parsed_data.get("name", "Name"), styles["Title"]))

    contact = parsed_data.get("contact", {})
    contact_line = " | ".join(filter(None, [
        contact.get("email", ""),
        contact.get("phone", ""),
        contact.get("location", "")
    ]))

    story.append(Paragraph(contact_line, styles["Normal"]))
    story.append(Spacer(1, 12))

    # Summary
    if parsed_data.get("summary"):
        story.append(Paragraph("SUMMARY", styles["Heading2"]))
        story.append(Paragraph(parsed_data["summary"], styles["Normal"]))

    doc.build(story)
    pdf_stream.seek(0)
    return pdf_stream


# ---------------------------
# DOCX GENERATION
# ---------------------------
def generate_optimized_resume_docx(original_text, missing_skills, analysis_data=None):
    doc = Document()
    doc.add_heading("Optimized Resume Suggestions", 1)

    doc.add_heading("Missing Skills", 2)
    for skill in missing_skills:
        doc.add_paragraph(skill, style="List Bullet")

    doc.add_heading("Original Resume", 2)
    doc.add_paragraph(original_text)

    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


# ---------------------------
# ROUTES
# ---------------------------
def register_resume_routes(app):

    from auth.auth import token_required


    # ---------------------------
    # MATCH RESUME
    # ---------------------------
    @app.route('/api/match', methods=['POST'])
    @token_required
    def process_match():

        can_use, message = check_usage_limit(g.user_id, 'analysis')
        if not can_use:
            return jsonify({"error": message}), 429

        if not model:
            return jsonify({"error": "LLM not available"}), 500

        if 'job_description' not in request.form:
            return jsonify({"error": "Missing job description"}), 400

        job_description = request.form['job_description']

        # Resume input
        if 'resume' in request.files:
            resume_file = request.files['resume']
            resume_text = extract_text_from_pdf(resume_file.stream)
        else:
            return jsonify({"error": "Missing resume"}), 400

        if not resume_text:
            return jsonify({"error": "Could not extract resume text"}), 400

        prompt = f"""
        You are an expert Applicant Tracking System (ATS) Analyst and Resume Optimization Specialist. Your job is to compare a RESUME against a JOB DESCRIPTION with comprehensive scoring and recommendations.

        --- RESUME TEXT ---
        {resume_text}

        --- JOB DESCRIPTION TEXT ---
        {job_description}

        Analyze the two texts and provide a comprehensive evaluation:

        1. **Calculate Detailed Scores (0-100%):**
        - "keyword_match_score": Percentage of critical job keywords present in resume
        - "skills_alignment_score": How well candidate's skills align with job requirements
        - "experience_relevance_score": How relevant work experience is to the position
        - "formatting_structure_score": How well-structured resume matches ATS expectations
        - "seniority_fit_score": Whether experience level matches position seniority
        - "overall_match_score": Weighted average of all scores

        2. **Identify Matched Skills:** List 5-10 key professional skills/technologies present in BOTH documents.

        3. **Identify Missing Skills:** List 5-10 key professional skills/technologies required by JOB DESCRIPTION but NOT found in RESUME.

        4. **Identify Weakly Represented Skills:** List 3-5 skills that appear in both documents but with weak representation in the resume (mentioned once or briefly).

        5. **Identify Overused Terms:** List any keywords/phrases that appear excessively in the resume (3+ times) that should be varied.

        6. **Keyword Gap Analysis:** For missing skills, suggest specific resume sections where each could be naturally integrated.

        7. **Generate Recommendation:** Write brief (max 3 sentences), actionable advice to improve resume for this specific job.

        8. **Add to Resume Suggestions:** Provide 3-5 specific bullet points or phrases the candidate could add to their resume to improve match.

        Return results STRICTLY as a single JSON object. Do not add any other text.
        Required keys: "keyword_match_score", "skills_alignment_score", "experience_relevance_score", "formatting_structure_score", "seniority_fit_score", "overall_match_score", "matched_skills", "missing_skills", "weakly_represented_skills", "overused_terms", "keyword_gap_analysis", "recommendation_text", "add_to_resume_suggestions"

        All scores must be integers 0-100.
        All list fields must be arrays of strings.
        keyword_gap_analysis must be an object mapping missing skill to suggested resume section (e.g., {{"Python": "Technical Skills section", "Docker": "Projects section"}}).
        """

        try:
            response_text = llm_call(prompt)
            json_string = response_text.strip().replace("```json", "").replace("```", "")

            result = json.loads(json_string)

            record_usage(g.user_id, 'analysis', {"job": True})

            return jsonify(result)

        except Exception as e:
            return jsonify({"error": str(e)}), 500


    # ---------------------------
    # SAVE RESUME
    # ---------------------------
    @app.route('/api/resumes/save', methods=['POST'])
    @token_required
    def save_resume():

        if 'resume' not in request.files:
            return jsonify({"error": "No resume provided"}), 400

        file = request.files['resume']
        text = extract_text_from_pdf(file.stream)

        db = get_db()
        cursor = db.cursor()

        cursor.execute(
            "SELECT COUNT(*) AS count FROM saved_resumes WHERE user_id = %s",
            (g.user['id'],)
        )

        count = cursor.fetchone()['count']
        limit = MAX_SAVED_RESUMES.get(g.user['subscription_type'], 1)

        if count >= limit:
            return jsonify({"error": "Resume limit reached"}), 400

        cursor.execute(
            """
            INSERT INTO saved_resumes (user_id, filename, resume_text)
            VALUES (%s, %s, %s)
            RETURNING id
            """,
            (g.user['id'], file.filename, text)
        )

        resume_id = cursor.fetchone()['id']
        db.commit()

        return jsonify({"id": resume_id, "message": "Saved"})


    # ---------------------------
    # GET RESUMES
    # ---------------------------
    @app.route('/api/resumes', methods=['GET'])
    @token_required
    def get_saved_resumes():
        """Get user's saved resumes."""
        db = get_db()
        cursor = db.cursor()

        cursor.execute('''
            SELECT id, filename, created_at 
            FROM saved_resumes 
            WHERE user_id = %s 
            ORDER BY created_at DESC
        ''', (g.user['id'],))

        resumes = cursor.fetchall()

        return jsonify({
            "resumes": [dict(row) for row in resumes]
        })


    # ---------------------------
    # DELETE RESUME
    # ---------------------------
    @app.route('/api/resumes/<int:resume_id>', methods=['DELETE'])
    @token_required
    def delete_resume(resume_id):

        db = get_db()
        cursor = db.cursor()

        cursor.execute(
            "DELETE FROM saved_resumes WHERE id = %s AND user_id = %s",
            (resume_id, g.user['id'])
        )

        db.commit()

        return jsonify({"message": "Deleted"})


    # ---------------------------
    # GENERATE DOCS
    # ---------------------------
    @app.route('/api/generate-cv', methods=['POST'])
    @token_required
    def generate_cv():

        data = request.json

        pdf = generate_standard_resume_pdf(data["original_resume_text"])

        return send_file(
            pdf,
            mimetype="application/pdf",
            as_attachment=True,
            download_name="cv.pdf"
        )